"""文件解析器（MVP：文本型 PDF / Excel / TXT；图片 OCR 由视觉模型在 M3 接入）。"""

from __future__ import annotations

from pathlib import Path

from pydantic import BaseModel, Field

try:
    import pypdf
except ImportError:  # pragma: no cover
    pypdf = None

try:
    import openpyxl
except ImportError:  # pragma: no cover
    openpyxl = None


class ParsedFile(BaseModel):
    file_id: str
    original_name: str
    mime_type: str = "application/octet-stream"
    text: str = ""
    table_summary: str = ""
    ocr_pending: bool = False
    image_path: str | None = None  # 图片本地路径，供视觉模型读取
    errors: list[str] = Field(default_factory=list)


class FileParser:
    MAX_TEXT_CHARS = 20_000
    BOM_HEADER_HINTS = ("名称", "图号", "数量", "材料", "型号", "规格")

    def parse(self, path: str | Path, file_id: str, original_name: str, mime_type: str) -> ParsedFile:
        path = Path(path)
        ext = path.suffix.lower()
        try:
            if ext == ".pdf":
                return self._parse_pdf(path, file_id, original_name, mime_type)
            if ext in (".xls", ".xlsx"):
                return self._parse_xlsx(path, file_id, original_name, mime_type)
            if ext in (".jpg", ".jpeg", ".png"):
                return self._parse_image(path, file_id, original_name, mime_type)
            return self._parse_text(path, file_id, original_name, mime_type)
        except Exception as exc:  # 单个文件失败不阻塞任务（AWF §6）
            return ParsedFile(
                file_id=file_id,
                original_name=original_name,
                mime_type=mime_type,
                errors=[f"{type(exc).__name__}: {exc}"],
            )

    def _parse_text(self, path: Path, file_id: str, original_name: str, mime_type: str) -> ParsedFile:
        text = path.read_text(encoding="utf-8", errors="replace")
        return ParsedFile(file_id=file_id, original_name=original_name, mime_type=mime_type, text=text[: self.MAX_TEXT_CHARS])

    def _parse_pdf(self, path: Path, file_id: str, original_name: str, mime_type: str) -> ParsedFile:
        if pypdf is None:
            return ParsedFile(file_id=file_id, original_name=original_name, mime_type=mime_type, errors=["pypdf not installed"])
        reader = pypdf.PdfReader(str(path))
        pages = []
        for i, page in enumerate(reader.pages, start=1):
            pages.append(f"[第{i}页]\n{page.extract_text() or ''}")
        return ParsedFile(
            file_id=file_id,
            original_name=original_name,
            mime_type=mime_type,
            text="\n\n".join(pages)[: self.MAX_TEXT_CHARS],
        )

    def _parse_xlsx(self, path: Path, file_id: str, original_name: str, mime_type: str) -> ParsedFile:
        if openpyxl is None:
            return ParsedFile(file_id=file_id, original_name=original_name, mime_type=mime_type, errors=["openpyxl not installed"])
        wb = openpyxl.load_workbook(str(path), data_only=True, read_only=True)
        lines = []
        bom_summary = []
        for sheet in wb.worksheets:
            rows = list(sheet.iter_rows(values_only=True))
            lines.append(f"工作表 {sheet.title}: {len(rows)} 行 x {sheet.max_column} 列")
            header = [str(c).strip() if c is not None else "" for c in rows[0]] if rows else []
            hit = [h for h in header if any(hint in h for hint in self.BOM_HEADER_HINTS)]
            if hit:
                bom_summary.append(f"{sheet.title}: 疑似BOM表，表头含 {', '.join(hit)}，数据行 {max(len(rows) - 1, 0)}")
            preview = rows[:20]
            for row in preview:
                vals = [str(v) if v is not None else "" for v in row]
                lines.append(" | ".join(vals))
        wb.close()
        return ParsedFile(
            file_id=file_id,
            original_name=original_name,
            mime_type=mime_type,
            text="\n".join(lines)[: self.MAX_TEXT_CHARS],
            table_summary="; ".join(bom_summary)[:2000],
        )

    def _parse_image(self, path: Path, file_id: str, original_name: str, mime_type: str) -> ParsedFile:
        return ParsedFile(
            file_id=file_id,
            original_name=original_name,
            mime_type=mime_type,
            text=f"[图片文件 {original_name}，由视觉模型读取]",
            ocr_pending=True,
            image_path=str(path),
        )
